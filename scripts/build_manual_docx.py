# -*- coding: utf-8 -*-
"""把 `docs/标准操作手册.md` 转成同目录的 `.docx`（给不懂技术的同事直接转发用）。

用法（在技能根目录）：

    python scripts/build_manual_docx.py                     # 转默认那份手册
    python scripts/build_manual_docx.py <输入.md> [-o <输出.docx>]

为什么要自己写，而不是装 pandoc：
    国内机器上 pandoc 常常装不起来（GitHub 不通、包管理器受限），而本仓库只需要
    「一份排版干净的中文手册」这一种用法 —— python-docx 是纯 pip 包、离线可装，
    规则又完全可控，比引入一个外部二进制更省事。

定位与局限（**不是**通用 markdown 引擎）：
    · 支持：H1~H3、段落、有序/无序/嵌套列表、表格、围栏代码块、引用块、分隔线、
      行内 `` `代码` ``、**加粗**、[链接](…)
    · 不支持：图片、脚注、嵌套表格、HTML。手册里也没用到这些。
    · 链接会被降级成纯文字：手册里的链接全是目录锚点（`#1-这套东西是什么`），
      docx 里没有对应目标，做成可点的反而会跳错地方。

这份 .docx 是**生成物**：要改内容请改 `.md`，再重跑本脚本（脚本会在文末写上这句）。
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    print("缺少 python-docx：pip install python-docx", file=sys.stderr)
    raise SystemExit(2)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "docs" / "标准操作手册.md"

CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)      # 标题用的深蓝
QUOTE_GRAY = RGBColor(0x59, 0x59, 0x59)

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


# ---------------------------------------------------------------- 小工具

def _set_style_font(style, name=CN_FONT, size=None, bold=None, color=None):
    """给样式设中英文字体。中文必须单独设 w:eastAsia，否则 Word 会回退到别的字体。"""
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def _shade(paragraph, fill="F2F2F2"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _left_bar(paragraph, color="BFBFBF", size="12"):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    bdr.append(left)
    ppr.append(bdr)


def _bottom_border(paragraph, color="BFBFBF", size="6"):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    ppr.append(bdr)


def _add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


# ---------------------------------------------------------------- 行内格式

def add_runs(paragraph, text, *, mono_size=None, color=None):
    """把一行 markdown 行内语法写进段落：**加粗** / `代码` / [文字](链接)"""
    for token in TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
            run.font.size = Pt(mono_size or 9.5)
            run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        elif token.startswith("[") and "](" in token:
            label = token[1:token.index("](")]
            run = paragraph.add_run(label)      # 锚点链接降级成纯文字
        else:
            run = paragraph.add_run(token)
        if color is not None:
            run.font.color.rgb = color
    return paragraph


# ---------------------------------------------------------------- 解析 md

def parse_markdown(text):
    """把 md 拆成块。够用就好：不建 AST，只认手册里出现过的元素。"""
    lines = text.replace("\r\n", "\n").split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", code))
            continue

        if not stripped:
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue

        if re.match(r"^-{3,}$", stripped) or re.match(r"^\*{3,}$", stripped):
            blocks.append(("hr",))
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue

        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", quote))
            continue

        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            ordered = bool(re.match(r"^\d+\.$", m.group(2)))
            blocks.append(("li", ordered, min(indent // 2, 3), m.group(3).strip()))
            i += 1
            continue

        para = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            n = nxt.strip()
            if (not n or n.startswith(("#", ">", "|", "```"))
                    or re.match(r"^(\s*)([-*]|\d+\.)\s+", nxt)
                    or re.match(r"^-{3,}$", n)):
                break
            para.append(n)
            i += 1
        blocks.append(("p", " ".join(para)))
    return blocks


# ---------------------------------------------------------------- 渲染

def setup_styles(doc):
    _set_style_font(doc.styles["Normal"], CN_FONT, 10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25

    _set_style_font(doc.styles["Heading 1"], CN_FONT, 19, True, ACCENT)
    _set_style_font(doc.styles["Heading 2"], CN_FONT, 14.5, True, ACCENT)
    _set_style_font(doc.styles["Heading 3"], CN_FONT, 12, True, RGBColor(0x33, 0x33, 0x33))
    for name, before, after in (("Heading 1", 18, 8), ("Heading 2", 14, 6), ("Heading 3", 10, 4)):
        pf = doc.styles[name].paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        if name in [s.name for s in doc.styles]:
            _set_style_font(doc.styles[name], CN_FONT, 10.5)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for attr, val in (("left_margin", 2.4), ("right_margin", 2.4),
                      ("top_margin", 2.2), ("bottom_margin", 2.2)):
        setattr(section, attr, Cm(val))

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("第 ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = QUOTE_GRAY
    _add_page_number(footer_p)
    tail = footer_p.add_run(" 页")
    tail.font.size = Pt(8.5)
    tail.font.color.rgb = QUOTE_GRAY


def render(doc, blocks):
    for block in blocks:
        kind = block[0]

        if kind == "h":
            _, level, text = block
            doc.add_heading(text, level=min(level, 3))

        elif kind == "p":
            add_runs(doc.add_paragraph(), block[1])

        elif kind == "li":
            _, ordered, level, text = block
            style = "List Number" if ordered else "List Bullet"
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph()
            add_runs(p, text)
            if level:
                p.paragraph_format.left_indent = Cm(0.75 + 0.6 * level)
            p.paragraph_format.space_after = Pt(2)

        elif kind == "code":
            lines = block[1] or [""]
            for idx, raw in enumerate(lines):
                p = doc.add_paragraph()
                run = p.add_run(raw if raw.strip() else " ")
                run.font.name = MONO_FONT
                run._element.rPr.rFonts.set(qn("w:eastAsia"), MONO_FONT)
                run.font.size = Pt(9)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(4 if idx == 0 else 0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Cm(0.4)
                _shade(p)

        elif kind == "quote":
            for idx, line in enumerate(block[1]):
                p = doc.add_paragraph()
                add_runs(p, line, color=QUOTE_GRAY)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(4 if idx == 0 else 0)
                _left_bar(p)

        elif kind == "table":
            rows = block[1]
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=cols)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                cells = table.add_row().cells
                for c_idx in range(cols):
                    value = row[c_idx] if c_idx < len(row) else ""
                    cell_p = cells[c_idx].paragraphs[0]
                    add_runs(cell_p, value, mono_size=9)
                    cell_p.paragraph_format.space_after = Pt(2)
                    for run in cell_p.runs:
                        run.font.size = Pt(9.5)
                        if r_idx == 0:
                            run.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

        elif kind == "hr":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            _bottom_border(p)


def main() -> int:
    ap = argparse.ArgumentParser(description="把手册 md 转成 docx（纯 python-docx）")
    ap.add_argument("input", nargs="?", default=str(DEFAULT_INPUT))
    ap.add_argument("-o", "--output", default=None)
    args = ap.parse_args()

    src = Path(args.input).resolve()
    if not src.exists():
        print(f"找不到输入文件：{src}", file=sys.stderr)
        return 1
    out = Path(args.output).resolve() if args.output else src.with_suffix(".docx")

    text = src.read_text(encoding="utf-8")
    blocks = parse_markdown(text)

    doc = Document()
    setup_styles(doc)
    render(doc, blocks)

    note = doc.add_paragraph()
    run = note.add_run(
        f"本文件由 scripts/build_manual_docx.py 从 {src.name} 自动生成 —— "
        "要改内容请改 .md 再重跑脚本，不要直接编辑本文件。"
    )
    run.font.size = Pt(8.5)
    run.font.color.rgb = QUOTE_GRAY
    _bottom_border(note)

    doc.save(out)
    headings = sum(1 for b in blocks if b[0] == "h")
    tables = sum(1 for b in blocks if b[0] == "table")
    codes = sum(1 for b in blocks if b[0] == "code")
    print(f"已生成：{out}")
    print(f"  段落块 {sum(1 for b in blocks if b[0] == 'p')} · 标题 {headings} · 表格 {tables} · 代码块 {codes}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
